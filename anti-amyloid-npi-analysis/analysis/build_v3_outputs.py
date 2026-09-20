# -*- coding: utf-8 -*-
"""網羅的探索の結果を1つのExcelにまとめ、レポートに感度解析節を追加する。"""
from __future__ import annotations

import csv
import io
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "output_v3")

SHEETS = [
    ("探索の診断", "v3_探索の診断.csv"),
    ("族別の内訳", "v3_families.csv"),
    ("全候補の検定結果", "v3_全候補の検定結果.csv"),
    ("ANCOVA", "v3_ANCOVA.csv"),
    ("全体検定", "v3_全体検定.csv"),
    ("群分けの感度解析", "v3_群分けの感度解析.csv"),
    ("定義に直結する項目", "v3_定義に直結する項目.csv"),
    ("min-p帰無分布", "v3_min-p帰無分布.csv"),
    ("症例別一覧", "v3_症例別一覧.csv"),
]


def read(name: str) -> list[list[str]]:
    p = os.path.join(OUT, name)
    if not os.path.exists(p):
        return []
    with io.open(p, encoding="utf-8-sig", newline="") as f:
        return [r for r in csv.reader(f)]


def main() -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("README")
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="44546A")
    for i, line in enumerate([
        ["NPIアパシースコア上昇群と非上昇群で差を示す項目の網羅的探索"],
        [""],
        ["この解析は既存データを確認したのちに行ったものであり、"
         "研究開始前に規定した解析ではない。"],
        ["有意になった項目だけを選び出す目的では作成していない。"
         "列挙した候補はすべて「全候補の検定結果」シートに収載している。"],
        [""],
        ["検定", "連続量は Mann-Whitney U の完全並べ替え検定（両側、同順位は中間順位）。"
         "2値変数は Fisher 正確確率検定（両側）。"],
        ["多重性", "Benjamini-Hochberg の FDR と、min-p 並べ替え検定"
         "（Westfall-Young）による FWER 補正p値を併記。"],
        ["全体検定", "エネルギー距離検定により、プロファイル全体の差を"
         "多重比較の影響を受けない1回の検定で評価。"],
        ["欠測", "欠測は0点として扱っていない。値の訂正・補完・除外は行っていない。"],
        [""],
        ["結論", "2,502項目を検定し、多重性を補正して有意となった項目はなかった。"
         "これは「差がない」という意味ではなく、"
         "この例数では差を検出できなかったという意味である。"],
    ], start=1):
        for j, v in enumerate(line, start=1):
            ws.cell(i, j, v)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=13)

    for title, name in SHEETS:
        rows = read(name)
        if not rows:
            continue
        s = wb.create_sheet(title[:31])
        for r in rows:
            s.append(r)
        for c in s[1]:
            c.font = hdr
            c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        s.freeze_panes = "A2"
        for j in range(1, s.max_column + 1):
            w = max((len(str(s.cell(i, j).value or "")) for i in
                     range(1, min(s.max_row, 200) + 1)), default=8)
            s.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 9), 60)
    path = os.path.join(OUT, "網羅的探索表_上昇群と非上昇群.xlsx")
    wb.save(path)
    print("saved", path, "sheets:", len(wb.sheetnames))

    # レポートへ感度解析節を追加
    rp = os.path.join(OUT, "v3_網羅的探索レポート.md")
    txt = io.open(rp, encoding="utf-8").read()
    rows = read("v3_群分けの感度解析.csv")
    if rows and "## 8. 群分けの切り方を変えた場合" not in txt:
        h = rows[0]
        keep = ["群分けの定義", "対象例数", "上昇群n", "対照群n", "検定した候補",
                "p<0.05に到達しうる候補", "最小p", "BH-FDR q<0.05",
                "FWER補正p（min-p並べ替え）", "備考"]
        idx = [h.index(k) for k in keep if k in h]
        add = ["\n## 8. 群分けの切り方を変えた場合\n",
               "主解析の切り方（ΔNPIアパシー≧1点）だけでなく、"
               "他の切り方でも同じ候補集合を検定した。"
               "**切り方を変えること自体が多重性であり、"
               "最小pが小さい切り方を選んで採用してはならない。**"
               "どの切り方も NPI アパシー得点から定義されるため、"
               "得点そのものに由来する候補（初回アパシーの有無）は"
               "同義反復として一律に除外している。\n",
               "| " + " | ".join(keep[i] if i < len(keep) else "" for i in
                                  range(len(idx))) + " |",
               "|" + "---|" * len(idx)]
        for r in rows[1:]:
            add.append("| " + " | ".join(r[i] if i < len(r) else "" for i in idx) + " |")
        add.append("\nいずれの切り方でも FWER 補正p値は有意水準に達しなかった。\n")
        io.open(rp, "w", encoding="utf-8").write(txt + "\n".join(add))
        print("report updated")


if __name__ == "__main__":
    main()
