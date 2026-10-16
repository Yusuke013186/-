# -*- coding: utf-8 -*-
"""初回NPI-Apathy score 1点以上の4例について、症例詳細と確認事項をWordにまとめる。

上野先生から「改善しなかった2例の詳細が必要」との依頼を受けて作成する補足資料。
値はすべて解析用ワークブックから読み込み、手入力しない。
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from npi_increase_analysis import build_cases, read_background, MMSE_SUB, CDR_DOM

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "abstract_v4", "症例詳細_初回アパシー4例と確認事項.docx")

MINCHO, GOTHIC = "ＭＳ 明朝", "ＭＳ ゴシック"


def font(run, name=MINCHO, size=10.5, bold=False, color=None):
    run.font.name = "Century"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para(doc, text, size=10.5, bold=False, name=MINCHO, after=4, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Mm(indent)
        p.paragraph_format.first_line_indent = Mm(-indent)
    font(p.add_run(text), name=name, size=size, bold=bold)
    return p


def head(doc, text):
    p = para(doc, text, size=12, bold=True, name=GOTHIC, after=4)
    p.paragraph_format.space_before = Pt(10)
    return p


def table(doc, header, rows, widths, bold_cells=frozenset()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        font(c.paragraphs[0].add_run(h), name=GOTHIC, size=9.5, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        sh = c._element.get_or_add_tcPr()
        shd = sh.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "E7E9EE"})
        sh.append(shd)
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            font(p.add_run(str(v)), size=9.5, bold=((ri, i) in bold_cells))
            if i > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Mm(w)
    return t


def g(x):
    return "－" if x is None else (f"{x:.1f}" if isinstance(x, float) and x % 1 else f"{x:g}")


def g1(x):
    return "－" if x is None else f"{x:.1f}"


def sg(x):
    return f"{x:+g}".replace("-", "−").replace("+0", "±0")


def main() -> None:
    cases, _ = build_cases()
    BG = read_background()
    inc = [c for c in cases if c.step == "S5_解析対象"]
    four = [c for c in inc if c.npi_bl.score >= 1]
    dec = sorted([c for c in four if c.npi_fu.score < c.npi_bl.score], key=lambda c: int(c.pid[1:]))
    non = sorted([c for c in four if c.npi_fu.score >= c.npi_bl.score],
                 key=lambda c: -c.npi_bl.score)          # 抄録の記載順（2→3点、1→3点）
    order = dec + non
    labels = [f"低下例1\n{dec[0].pid}", f"低下例2\n{dec[1].pid}",
              f"非低下例1\n{non[0].pid}", f"非低下例2\n{non[1].pid}"]

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Mm(18))

    p = para(doc, "初回NPI-Apathy score 1点以上の4例　症例詳細と確認事項", size=14, bold=True,
             name=GOTHIC, after=2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(doc, "総会抄録（上野先生修正版）の補足資料　／　抄録本文には含めません", size=9.5, after=8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    para(doc, f"解析対象14例（全例レカネマブ）のうち、初回NPI-Apathy scoreが1点以上であったのは4例で、"
              f"いずれも女性であった。初回0点の10例はスコアが低下しえないため、"
              f"「初回1点以上」の規則で比較対象が機械的に決まる。低下例2例（{dec[0].pid}、{dec[1].pid}）と"
              f"非低下例2例（{non[0].pid}、{non[1].pid}）の値はすべて解析用ワークブックから読み込んだものである。"
              f"研究番号は匿名IDである。")

    # ---------------------------------------------------------- 表1
    head(doc, "表1　背景とNPI-Apathy score")
    def npi(r):
        return f"{g(r.score)}（{g(r.freq)}×{g(r.sev)}）"
    def ndate(r):
        return str(r.date).replace("-", "/") if r.date else f"{r.ym[0]}/{r.ym[1]}（日未記載）"
    rows = [
        ["年齢・性別"] + [f"{BG[c.pid]['年齢']:g}歳・{BG[c.pid]['性別']}" for c in order],
        ["投与開始日"] + [BG[c.pid]["投与開始日"].strftime("%Y/%m/%d") for c in order],
        ["NPI-Apathy 初回\n（頻度×重症度）"] + [npi(c.npi_bl) for c in order],
        ["NPI-Apathy 追跡時\n（頻度×重症度）"] + [npi(c.npi_fu) + ("※1" if c.npi_fu.rule_imputed else "")
                                         for c in order],
        ["アパシー有無欄\n初回→追跡時"] + [f"{c.npi_bl.presence}→{c.npi_fu.presence}" for c in order],
        ["NPI評価日 初回"] + [ndate(c.npi_bl) for c in order],
        ["NPI評価日 追跡時\n（原資料の時点）"] + [f"{ndate(c.npi_fu)}\n（{c.npi_fu.orig_label}）" for c in order],
        ["易怒性\n初回→追跡時"] + [f"{c.npi_bl.irrit}→{c.npi_fu.irrit}" for c in order],
        ["認知機能の追跡時点\n（評価日）"] + [
            f"{c.followup_visit('MMSE').label}\n（{str(c.followup_visit('MMSE').date).replace('-', '/') if c.followup_visit('MMSE').date else '評価日記載なし'}）"
            for c in order],
    ]
    table(doc, ["項目"] + labels, rows, [38, 34, 34, 34, 34])
    para(doc, "※1　追跡時は有無欄が「なし」で頻度・重症度欄が空欄のため、研究上の記録ルールにより0点とした"
              "（NPIでは主質問が「なし」のとき当該項目は0点となる）。", size=9, after=2)

    # ---------------------------------------------------------- 表2
    head(doc, "表2　MMSE（初回→追跡時）　太字は変化した項目")
    rows, bold = [], set()
    keys = [("合計", "__total__")] + [(k, k) for k in MMSE_SUB]
    for i, (lab, k) in enumerate(keys):
        vals = []
        for j, c in enumerate(order, 1):
            b, f = c.bl_fu("MMSE", k)
            vals.append(f"{g(b)}→{g(f)}")
            if b != f:
                bold.add((i, j))
        rows.append([lab] + vals)
    sums = []
    for c in order:
        d = sum(c.bl_fu("MMSE", k)[1] - c.bl_fu("MMSE", k)[0] for k in MMSE_SUB)
        tot = c.delta("MMSE", "__total__")
        sums.append(f"{sg(d)}／{sg(tot)}{'（一致）' if d == tot else '（不一致）'}")
    rows.append(["下位項目Δの総和／合計Δ"] + sums)
    table(doc, ["項目"] + labels, rows, [38, 34, 34, 34, 34], bold_cells=frozenset(bold))

    # ---------------------------------------------------------- 表3
    head(doc, "表3　CDR（初回→追跡時）　太字は変化した項目")
    rows, bold = [], set()
    keys = [("CDR-SB", "__total__"), ("Global CDR", "__global__")] + [(k, k) for k in CDR_DOM]
    for i, (lab, k) in enumerate(keys):
        vals = []
        for j, c in enumerate(order, 1):
            b, f = c.bl_fu("CDR", k)
            vals.append(f"{g1(b)}→{g1(f)}")
            if b != f:
                bold.add((i, j))
        rows.append([lab] + vals)
    table(doc, ["項目"] + labels, rows, [38, 34, 34, 34, 34], bold_cells=frozenset(bold))

    # ---------------------------------------------------------- 非低下例の要約
    head(doc, "非低下例2例の経過（上野先生ご依頼の詳細）")

    def changed_items(c, sheet, keys_):
        out = []
        for k in keys_:
            b, f = c.bl_fu(sheet, k)
            if b != f:
                out.append(f"{k}{g1(b) if sheet == 'CDR' else g(b)}→{g1(f) if sheet == 'CDR' else g(f)}")
        return "、".join(out) if out else "なし"

    for i, c in enumerate(non, 1):
        bt, ft = c.bl_fu("MMSE", "__total__")
        bs, fs = c.bl_fu("CDR", "__total__")
        tb, tf = c.bl_fu("MMSE", "時間見当識")
        extra = "初回の時間見当識は満点（5点）で、上昇の余地はなかった。" if tb == 5 else \
            f"初回の時間見当識は{g(tb)}点で上昇の余地があったが、{g(tf)}点となった。"
        para(doc, f"非低下例{i}（{c.pid}、{BG[c.pid]['年齢']:g}歳・{BG[c.pid]['性別']}）："
                  f"NPI-Apathy score {g(c.npi_bl.score)}→{g(c.npi_fu.score)}点"
                  f"（頻度{g(c.npi_bl.freq)}→{g(c.npi_fu.freq)}、重症度{g(c.npi_bl.sev)}→{g(c.npi_fu.sev)}）。"
                  f"MMSE {g(bt)}→{g(ft)}点（変化した下位項目：{changed_items(c, 'MMSE', MMSE_SUB)}）、"
                  f"CDR-SB {g1(bs)}→{g1(fs)}点（変化した領域：{changed_items(c, 'CDR', CDR_DOM)}）。"
                  f"{extra}易怒性は{c.npi_bl.irrit}→{c.npi_fu.irrit}。"
                  f"認知機能の追跡は{c.followup_visit('MMSE').label}時点。", indent=0)

    # ---------------------------------------------------------- 確認事項
    head(doc, "確認をお願いしたい事項")
    l60 = next(c for c in four if c.pid == "L60")
    items = [
        f"【重要】L60（非低下例2）の初回NPIの評価日は原資料に「2024.6.（日未記載）」と記載されている。"
        f"投与開始日（{BG['L60']['投与開始日'].strftime('%Y/%m/%d')}）および認知機能の初回評価"
        f"（{str(l60.baseline_visit('MMSE').date).replace('-', '/')}）より約14か月前であり、"
        f"追跡時の原資料の時点「12か月」（2026/8/17）とも整合しない。ワークブックにも「対応を要確認」と"
        f"記載がある。比較対象は2例しかないため、この1例の初回値の扱いは結果に直結する。"
        f"値は変更せず原記載のまま用いた。原票の確認をお願いしたい。",
        "L4（低下例1）の初回アパシー有無欄は「判定保留」（有・無の双方に丸と訂正跡）。"
        "頻度1×重症度1＝1点は原資料に記載されている。低下（1→0点）の判定はこの1点に依存する。",
        "L4・L59（低下例2例）の追跡時0点は、有無欄「なし」で頻度・重症度が空欄のため、"
        "研究上の記録ルールにより0点としたもの（表1※1）。",
        "認知機能の追跡はNPIの追跡より約6か月早い。L58・L59・L60は認知機能6か月（2026年2月）に対し"
        "NPIは原資料12か月（2026年8月）、L4は認知機能24か月に対しNPIは原資料30か月。"
        "NPIの追跡時点に対応する認知機能評価が存在しないためで、抄録には「評価間隔はNPI-Apathy scoreが"
        "12～30か月で，認知機能はNPI追跡評価日に最も近い時点の値を用いた（6～24か月）」と書き分けた。",
        "非低下例1（L58）は初回の時間見当識が満点（5点）で、上昇の余地がない。"
        "非低下例2（L60）は初回3点で上昇の余地があったが2点となった。"
        "抄録の結論ではこの初回値の違いを限界として記載した。",
        "易怒性は低下2例と非低下例2（L60）で新規に出現し、非低下例1（L58）は初回からあった。"
        "4例で見ると新規出現は低下例に限られないため、堀内版にあった「2例とも易怒性が新規に出現した」は"
        "上野先生版と同じく本文に含めていない。",
        "「後方視的に解析した」は上野先生版の記載をそのまま用いた。研究デザインはデータからは確認できない。",
        "対象期間「2024年1月～2026年7月」は旧稿からの記載。患者背景ファイルで投与開始日が記録されているのは"
        "97件中63件で、範囲は2024/1/13～2026/2/18であり、2026年7月までは確認できない。",
        "MMSE合計点と下位項目合計は、対象14例のすべての評価時点で一致した（不一致0件、表2最下行）。"
        "ただし添付PDFの原票との照合は行えていない。",
    ]
    for i, t in enumerate(items, 1):
        para(doc, f"{i}．{t}", size=10, indent=5)

    # ---------------------------------------------------------- 変更点
    head(doc, "上野先生版から変更・補充した箇所")
    ch = [
        "空欄の補充：不変7例・上昇5例、年齢71～81歳・女性4例、非低下2例のスコア（2→3点、1→3点）、"
        "非低下2例の時間見当識／MMSE／CDR-SB。",
        "【方法】「評価間隔は6～24か月」→「評価間隔はNPI-Apathy scoreが12～30か月で，認知機能はNPI追跡評価日に"
        "最も近い時点の値を用いた（6～24か月）」。6～24か月は認知機能の間隔で、NPIの間隔（原資料の時点で"
        "12～30か月）とは異なるため。後半は堀内版で採用した記載。",
        "【方法】堀内版から「（頻度×重症度，0～12点）」「（全例レカネマブ）」を補った。",
        "【方法】「初回NPI-Apathy score　1点以上」の全角スペースを半角に統一した。",
        "【結果】「不変7例」に「（いずれも0点のまま）」を補った。初回1点以上が4例に限られる理由が読み取れるため。",
        "【結論】堀内版の床効果・平均への回帰の記載を、4例の値に合わせて1文補った"
        "（初回時間見当識：低下2例がいずれも2点、非低下2例が5点と3点）。",
        "テンプレに元から記載されていた低下例2例の値（2→3／23→23／2.0→4.0、2→4／22→25／1.5→1.5）は"
        "元データと一致することを確認した。",
    ]
    for i, t in enumerate(ch, 1):
        para(doc, f"{i}．{t}", size=10, indent=5)

    head(doc, "堀内版にあり、本文に含めなかった記載（必要なら追加できます）")
    nc = [
        "示唆の一文：「これらの限界はあるが，本所見はアパシーの改善と時間見当識の改善が関連して推移する"
        "可能性を示唆する．」（上野先生版の結論は「両者の変化量の関連を検証する必要がある」で結ばれているため）",
        "「全例が治療例であり薬剤の効果は論じられない．」",
        "「群分けは既存データ確認後に設定した．」（題名に「探索的検討」とあり、事前規定の解析とは記載していない）",
    ]
    for i, t in enumerate(nc, 1):
        para(doc, f"{i}．{t}", size=10, indent=5)

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
