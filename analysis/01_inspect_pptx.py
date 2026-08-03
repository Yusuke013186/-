# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu
p = Presentation("8月5日に見せるスライド.pptx")
print("slide size:", p.slide_width, p.slide_height, "=", p.slide_width/914400, "x", p.slide_height/914400, "inch")
print("layouts:", [(i, l.name) for i, l in enumerate(p.slide_masters[0].slide_layouts)])
for si, s in enumerate(p.slides):
    print(f"\n===== SLIDE {si+1}  layout={s.slide_layout.name} =====")
    for sh in s.shapes:
        print(f"  [{sh.shape_type}] name={sh.name!r} pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})")
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                runs = [(r.text, r.font.size.pt if r.font.size else None, r.font.name, r.font.bold,
                         r.font.color.rgb if r.font.color and r.font.color.type is not None else None) for r in para.runs]
                if runs: print("      P:", runs)
        if sh.has_table:
            t = sh.table
            print(f"      TABLE {len(t.rows)}x{len(t.columns)}")
            for ri,row in enumerate(t.rows):
                print("       ", [c.text for c in row.cells])
        if sh.has_chart:
            ch = sh.chart
            print("      CHART type:", ch.chart_type)
            try:
                print("      cats:", list(ch.plots[0].categories))
            except Exception as e: print("      cats err", e)
            for ser in ch.series:
                print("      series:", ser.name, list(ser.values))
    if s.has_notes_slide:
        print("  NOTES:", s.notes_slide.notes_text_frame.text[:1500])
